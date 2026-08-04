import tempfile
import os
from io import BytesIO

from locator.data_center.batch_register import BatchRegisterLocator
from locator.data_center.data_center import DataCenterLocator
from locator.InCMS import InCMSLocator
from object.basepage import BasePage
from object.data_center.data_center import DataCenterPage
from utils.log import logger

import allure
import openpyxl


class BatchRegisterPage(BasePage):
    """批量注册页"""

    def __init__(self, page):
        super().__init__(page)

    @allure.step("进入批量注册界面")
    def enter_batch_register(self):
        """从数据中心进入批量注册界面"""
        logger.debug("进入批量注册界面")
        self.click(InCMSLocator.register_button)
        self.page.wait_for_timeout(500)
        self.click(InCMSLocator.batch_register)
        logger.info("已进入批量注册界面")

    @allure.step("下载批量注册模板")
    def download_register_template(self):
        """下载批量注册 Excel 模板，返回 openpyxl Workbook"""
        logger.debug("开始下载批量注册模板")
        frame = self.get_frame(BatchRegisterLocator.batch_register_frame)
        with self.page.expect_download() as download_info:
            frame.click(BatchRegisterLocator.download_template_btn)
        download = download_info.value
        logger.debug(f"模板文件名: {download.suggested_filename}")
        data = download.path().read_bytes()
        logger.debug(f"模板文件大小: {len(data)} bytes")
        assert len(data) > 0, "批量注册模板文件为空"
        wb = openpyxl.load_workbook(BytesIO(data))
        logger.info(f"批量注册模板下载成功，共 {len(wb.sheetnames)} 个 sheet")
        return wb

    @staticmethod
    @allure.step("生成批量注册 Excel")
    def generate_register_excel(headers, fields, batch_number, smiles, register_amount):
        """
        用预期字典生成带数据的注册 Excel 文件
        :param headers: 第1行字段ID列表
        :param fields: 第2行字段名列表
        :param batch_number: 批号
        :param smiles: SMILES
        :param register_amount: 注册的量
        :return: 文件路径
        """
        logger.debug("开始生成批量注册 Excel")
        wb = openpyxl.Workbook()
        ws = wb.active

        for col_idx, header in enumerate(headers, 1):
            ws.cell(row=1, column=col_idx, value=header)
        for col_idx, field in enumerate(fields, 1):
            ws.cell(row=2, column=col_idx, value=field)

        # 字段名 → 列索引 → 填入数据
        col_map = {f: idx + 1 for idx, f in enumerate(fields)}

        data_map = {
            "*批号(带星号的为必填列)": batch_number,
            "SMILES/Mol": smiles,
            "*注册的量[数字+体积或质量单位(L/ml/g/mg等)]": register_amount,
        }

        for field_name, value in data_map.items():
            col = col_map.get(field_name)
            if col:
                ws.cell(row=3, column=col, value=value)
                logger.debug(f"填入字段 '{field_name}' 于第3行第{col}列: {value}")
            else:
                logger.debug(f"未找到字段 '{field_name}'")

        save_path = os.path.join(tempfile.gettempdir(), f"batch_register_{batch_number}.xlsx")
        wb.save(save_path)
        logger.info(f"批量注册 Excel 已生成: {save_path}")
        return save_path



    @allure.step("设置自动提交，关闭自动登记样品和入库")
    def setup_auto_submit_and_register(self):
        """设置：上传成功后自动提交注册=开，自动登记样品=关，自动入库=关"""
        logger.debug("开始设置自动提交与登记选项")
        frame = self.get_frame(BatchRegisterLocator.batch_register_frame)

        # 自动提交注册：确保开启
        loc = frame.find(BatchRegisterLocator.auto_submit_btn)
        if not loc.evaluate("el => el.classList.contains('on')"):
            logger.debug("自动提交注册未开启，点击开启")
            frame.click(BatchRegisterLocator.auto_submit_btn)
            self.page.wait_for_timeout(1000)
            assert loc.evaluate("el => el.classList.contains('on')"), "自动提交注册未开启"
        logger.info("自动提交注册已开启")

        # 自动登记样品：确保关闭
        loc = frame.find(BatchRegisterLocator.auto_register_sample_btn)
        if loc.evaluate("el => el.classList.contains('on')"):
            logger.debug("自动登记样品已开启，点击关闭")
            frame.click(BatchRegisterLocator.auto_register_sample_btn)
            self.page.wait_for_timeout(1000)
            assert not loc.evaluate("el => el.classList.contains('on')"), "自动登记样品未关闭"
        logger.info("自动登记样品已关闭")

        # 自动入库：确保关闭
        loc = frame.find(BatchRegisterLocator.auto_inbound_btn)
        if loc.evaluate("el => el.classList.contains('on')"):
            logger.debug("自动入库已开启，点击关闭")
            frame.click(BatchRegisterLocator.auto_inbound_btn)
            self.page.wait_for_timeout(1000)
            assert not loc.evaluate("el => el.classList.contains('on')"), "自动入库未关闭"
        logger.info("自动入库已关闭")

        logger.info("自动提交与登记选项设置完成")

    @allure.step("上传批量注册Excel")
    def upload_register_excel(self, file_path: str):
        """上传 Excel 文件并点击上传按钮"""
        logger.debug(f"上传批量注册文件: {file_path}")
        frame = self.get_frame(BatchRegisterLocator.batch_register_frame)
        frame.input_file(BatchRegisterLocator.excel_upload_input, file_path)
        logger.debug("文件已选择，等待上传按钮可用")
        self.page.wait_for_timeout(1000)
        frame.click(BatchRegisterLocator.upload_btn)
        os.remove(file_path)
        logger.info("批量注册文件上传完成，临时文件已清理")

    @allure.step("等待下载任务完成")
    def wait_for_download_complete(self, timeout=20000):
        """等待 floating_inbox 出现后消失"""
        logger.debug("等待下载框出现")
        self.wait_for(DataCenterLocator.floating_inbox, "visible", timeout=timeout)
        logger.debug("等待下载框消失")
        self.wait_for(DataCenterLocator.floating_inbox, "hidden", timeout=timeout)
        self.page.wait_for_timeout(3000)
        logger.info("下载任务已完成")

    @allure.step("在物质列表中查找批号")
    def assert_material_exists(self, batch_number):
        """调用接口刷新物质列表，断言指定批号存在"""
        logger.debug(f"在物质列表中查找批号: {batch_number}")
        dc = DataCenterPage(self.page)
        api_data = dc.get_material_list_from_api()
        data_list = api_data["data"]["data"]["dataList"]
        found = any(
            dc._extract_text(item.get("expCode")) == batch_number
            for item in data_list
        )
        assert found, f"未在物质列表中找到批号: {batch_number}"
        logger.info(f"物质列表中已找到批号: {batch_number}")
