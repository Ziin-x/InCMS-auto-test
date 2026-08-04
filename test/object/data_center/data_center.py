from io import BytesIO
from datetime import datetime, timezone, timedelta

from config.Expection import ElementNotFoundError
from config.setting import MaterialListAPI
from locator.InCMS import InCMSLocator
from locator.data_center.data_center import DataCenterLocator
from object.basepage import BasePage
from utils.load_ele_param import load_ele_param
from utils.log import logger

import allure
import openpyxl


class DataCenterPage(BasePage):
    def __init__(self,page):
        super().__init__(page)

    @allure.step("进入数据中心-全部项目界面")
    def enter_all_project(self):
        logger.debug("进入数据中心-全部项目界面")
        self.ensure_dropdown_expanded(InCMSLocator.data_center, InCMSLocator.all_projects)
        self.click(InCMSLocator.all_projects)
        logger.info("已进入数据中心-全部项目界面")

    @allure.step("断言项目在数据中心可见")
    def assert_project_visible(self,project_name):
        logger.debug(f"断言项目在数据中心可见: {project_name}")
        self.ensure_dropdown_expanded(InCMSLocator.data_center, InCMSLocator.all_projects)
        project_loc = load_ele_param(InCMSLocator.data_center_project, project_name)
        try:
            self.wait_for(project_loc,"visible")
            logger.info(f"项目 '{project_name}' 在数据中心可见")
        except Exception as e:
            logger.error(f"断言项目在数据中心不可见，错误：{e}")
            raise ElementNotFoundError(
            f"项目 '{project_name}' 在数据中心的下拉列表中不可见，等待超时或元素不存在。原始错误：{e}"
        )

### ---------------------------- 项目详细界面 ---------------------------- ###

    def _is_tab_active(self, locator):
        loc = self.find(locator)
        return loc.evaluate("el => el.classList.contains('active')")

    @allure.step("切换到{tab_name}视图")
    def switch_view(self, tab_locator, tab_name):
        logger.debug(f"切换到{tab_name}视图")
        if self._is_tab_active(tab_locator):
            logger.info(f"已在{tab_name}视图")
            return
        logger.debug(f"当前不在{tab_name}视图，点击切换")
        self.click(tab_locator)
        self.page.wait_for_timeout(3000)
        if not self._is_tab_active(tab_locator):
            logger.error(f"切换到{tab_name}视图失败")
            raise RuntimeError(f"切换到{tab_name}视图失败")
        logger.info(f"已切换到{tab_name}视图")

    def _export_and_download(self, export_locator, file_type):
        """通用导出流程：点击更多→点击导出菜单→等待→下载框下载"""
        logger.debug(f"开始导出全部数据的 {file_type}")
        logger.debug("点击更多操作按钮")
        self.click(DataCenterLocator.more_button)
        logger.debug(f"点击全部导出{file_type}")
        self.click(export_locator)
        logger.debug("导出任务已提交，等待任务完成")
        self.page.wait_for_timeout(3000)

        logger.debug("点击下载框")
        self.click(DataCenterLocator.floating_inbox)
        logger.debug("点击下载任务结果按钮")
        with self.page.expect_download() as download_info:
            self.click_nth(DataCenterLocator.floating_inbox_download, 0)
        download = download_info.value
        logger.debug(f"下载文件名: {download.suggested_filename}")
        data = download.path().read_bytes()
        logger.debug(f"下载文件大小: {len(data)} bytes")
        return download, data

    @allure.step("导出全部数据的excel")
    def export_all_excel(self):
        try:
            download, data = self._export_and_download(DataCenterLocator.export_all_excel, "Excel")
            wb = openpyxl.load_workbook(BytesIO(data))
            logger.info(f"Excel 读取成功，共 {len(wb.sheetnames)} 个 sheet")
            return wb
        except Exception as e:
            logger.error(f"导出 Excel 失败: {e}")
            raise

    @allure.step("导出全部数据的SDF")
    def export_all_sdf(self):
        try:
            download, data = self._export_and_download(DataCenterLocator.export_all_sdf, "SDF")
            assert len(data) > 0, "SDF 文件为空"
            logger.info(f"SDF 导出成功，文件大小: {len(data)} bytes")
            return data
        except Exception as e:
            logger.error(f"导出 SDF 失败: {e}")
            raise

    @allure.step("获取当前项目物质列表接口数据")
    def get_material_list_from_api(self):
        logger.debug("开始监听物质列表接口")
        try:
            with self.page.expect_response(lambda r: MaterialListAPI in r.url) as response_info:
                logger.debug("刷新页面以触发接口请求")
                self.reload()
                logger.debug("已刷新页面，等待接口响应")
            response = response_info.value
            logger.debug(f"接口响应状态: {response.status}")
            data = response.json()
            item_count = len(data.get("data", {}).get("data", {}).get("dataList", []))
            logger.info(f"物质列表接口响应成功，共 {item_count} 条记录")
            return data
        except Exception as e:
            logger.error(f"获取物质列表接口数据失败: {e}")
            raise

    @allure.step("比对 Excel 与接口数据")
    def compare_material_data(self, excel_wb, api_data):
        logger.debug("开始比对 Excel 与接口数据")
        try:
            api_list = api_data["data"]["data"]["dataList"]
            field_list = api_data["data"]["data"]["fieldList"]
            extra = api_data["data"]["extra"]
        except (KeyError, TypeError) as e:
            logger.error(f"解析 API 数据结构失败: {e}")
            raise

        ws = excel_wb.active
        logger.debug(f"Excel sheet: {ws.title}, 行数: {ws.max_row}, 列数: {ws.max_column}")

        headers = [cell.value for cell in ws[1]]
        logger.debug(f"Excel 表头共 {len(headers)} 列")

        # 构建 excel 行索引：按批号匹配，避免依赖排序
        try:
            exp_code_col = headers.index("批号") + 1
        except ValueError:
            logger.error("Excel 表头中未找到'批号'列")
            raise
        excel_map = {}
        for r in range(2, ws.max_row + 1):
            k = ws.cell(row=r, column=exp_code_col).value
            excel_map[k] = r
        logger.debug(f"Excel 行索引构建完成，共 {len(excel_map)} 行")

        # 构建查找表
        id_to_name = self._build_id_name_map(field_list)
        logger.debug(f"ID→名称映射表构建完成，共 {len(id_to_name)} 个字段")
        user_map = {u["id"]: u["name"] for u in extra.get("userInfo", [])}
        logger.debug(f"用户映射表构建完成，共 {len(user_map)} 个用户")

        mappings = self._get_field_mappings()
        logger.debug(f"比对字段共 {len(mappings)} 个")

        diffs = []
        for item in api_list:
            api_exp = self._extract_text(item.get("expCode")) or ""
            row_idx = excel_map.get(api_exp)
            if row_idx is None:
                logger.error(f"Excel 中未找到批号: {api_exp}")
                diffs.append({"exp_code": api_exp, "error": "Excel 中无对应行"})
                continue

            for mapping in mappings:
                try:
                    api_val = self._resolve_api_value(item, mapping, id_to_name, user_map)
                    excel_val = self._get_excel_cell(ws, row_idx, headers, mapping["excel_col"])
                    api_str = self._normalize_value(api_val)
                    excel_str = self._normalize_value(excel_val)
                    if api_str != excel_str:
                        logger.debug(f"差异: 批号={api_exp} 字段={mapping['api_field']} "
                                     f"API='{api_val}' Excel='{excel_val}'")
                        diffs.append({
                            "exp_code": api_exp,
                            "field": mapping["api_field"],
                            "api_value": api_val,
                            "excel_value": excel_val,
                        })
                except Exception as e:
                    logger.error(f"比对字段 {mapping['api_field']} 时出错: {e}")
                    raise

        if diffs:
            logger.error(f"比对不一致，共 {len(diffs)} 处差异")
            for d in diffs:
                logger.error(f"  批号={d.get('exp_code')} 字段={d.get('field')} "
                             f"API={d.get('api_value')} Excel={d.get('excel_value')}")
        else:
            logger.info("Excel 与接口数据比对一致，共 {} 条记录".format(len(api_list)))
        return diffs

    # ---------- 映射配置 ----------

    _MAPPINGS = None

    def _get_field_mappings(self):
        if self._MAPPINGS is not None:
            return self._MAPPINGS
        self._MAPPINGS = [
            {"api_field": "incmsCode",    "excel_col": "注册编号",     "fmt": "text"},
            {"api_field": "projectName",  "excel_col": "项目名称",     "fmt": "id_to_name"},
            {"api_field": "serialNumber", "excel_col": "流水号",       "fmt": "text"},
            {"api_field": "ics",          "excel_col": "ICS",          "fmt": "text"},
            {"api_field": "cas",          "excel_col": "CAS",          "fmt": "text"},
            {"api_field": "iupac",        "excel_col": "英文IUPAC",    "fmt": "text"},
            {"api_field": "iupacCn",      "excel_col": "中文IUPAC",    "fmt": "text"},
            {"api_field": "smiles",       "excel_col": "SMILES",       "fmt": "text"},
            {"api_field": "formula",      "excel_col": "分子式",        "fmt": "text"},
        ]
        logger.debug(f"字段映射初始化完成，共 {len(self._MAPPINGS)} 个字段")
        return self._MAPPINGS

    # ---------- 辅助方法 ----------

    def _build_id_name_map(self, field_list):
        id_map = {}
        for f in field_list:
            props = f.get("property", {}) or {}
            options = props.get("options", [])
            if options:
                id_map[f["fieldName"]] = {o["id"]: o["label"] for o in options}
        return id_map

    def _extract_text(self, val):
        if isinstance(val, list) and len(val) > 0 and isinstance(val[0], dict):
            return val[0].get("text", None)
        if isinstance(val, dict) and "text" in val:
            return val.get("text", None)
        return val

    def _resolve_api_value(self, item, mapping, id_map, user_map):
        field = mapping["api_field"]
        fmt = mapping["fmt"]
        raw = item.get(field)

        try:
            if fmt == "text":
                return self._extract_text(raw)
            elif fmt == "text_array":
                if isinstance(raw, list):
                    return ", ".join(d.get("text", "") for d in raw if isinstance(d, dict))
                return raw
            elif fmt == "text_nullable":
                return self._extract_text(raw)
            elif fmt in ("number", "number_nullable", "string_nullable"):
                return raw
            elif fmt == "id_to_name":
                name_map = id_map.get(field, {})
                if raw is not None:
                    result = name_map.get(str(raw), str(raw))
                    if result == str(raw):
                        logger.debug(f"字段 {field} 的值 '{raw}' 在选项列表中未找到对应名称")
                    return result
                return None
            elif fmt == "id_to_name_nullable":
                if raw is None:
                    return None
                name_map = id_map.get(field, {})
                return name_map.get(str(raw), str(raw))
            elif fmt == "user_id":
                return user_map.get(str(raw), str(raw)) if raw is not None else None
            elif fmt == "mass":
                if isinstance(raw, dict):
                    num = raw.get("number", "")
                    unit = raw.get("unitId", "")
                    return f"{num}{unit}" if num else None
                return raw
            elif fmt == "timestamp":
                if isinstance(raw, (int, float)):
                    return datetime.fromtimestamp(raw / 1000, tz=timezone.utc) \
                        .astimezone(timezone(timedelta(hours=8))) \
                        .strftime("%Y-%m-%d %H:%M:%S")
                return raw
            elif fmt == "empty_array":
                if isinstance(raw, list) and len(raw) == 0:
                    return None
                return raw
            elif fmt == "status_value":
                if isinstance(raw, dict) and raw.get("status") == "success":
                    vals = raw.get("value", [])
                    return None if len(vals) == 0 else vals
                return raw
            logger.debug(f"未知的 fmt 类型: {fmt}，字段: {field}")
            return raw
        except Exception as e:
            logger.error(f"解析字段 {field} (fmt={fmt}) 失败: {e}, raw={raw}")
            raise

    def _normalize_value(self, val):
        if val is None:
            return ""
        if isinstance(val, float):
            return f"{val:.6g}"
        if isinstance(val, bool):
            return str(val)
        return str(val).strip()

    def _get_excel_cell(self, ws, row_idx, headers, col_name):
        try:
            col_idx = headers.index(col_name)
            return ws.cell(row=row_idx, column=col_idx + 1).value
        except ValueError:
            logger.debug(f"Excel 中未找到列: {col_name}")
            return None
