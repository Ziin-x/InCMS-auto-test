import tempfile
import os

from locator.sample_center.batch_import import BatchImportLocator
from locator.sample_center.sample_list import SampleListLocator
from object.basepage import BasePage
from utils.log import logger

import allure
import openpyxl


class BatchImportPage(BasePage):
    """批量导入页"""

    def __init__(self, page):
        super().__init__(page)

    @allure.step("批量登记样品")
    def register_samples(self, file_path):
        """进入批量登记界面并上传 Excel 文件"""
        self.click(SampleListLocator.batch)
        self.click(SampleListLocator.batch_register)
        self.page.wait_for_timeout(2000)
        logger.debug("已进入批量登记界面")

        frame = self.get_frame(BatchImportLocator.batch_import_frame)
        # 点击菜单 → 文件
        frame.click(BatchImportLocator.intable_menu_btn)
        logger.debug("已点击菜单")
        frame.click(BatchImportLocator.file_menu_btn)
        logger.debug("已点击文件")

        # 上传 Excel
        frame.input_file(BatchImportLocator.import_excel_input, file_path)
        os.remove(file_path)
        logger.info("已上传样品登记文件，临时文件已清理")
        frame.click(BatchImportLocator.confirm_btn)
        logger.debug("已点击确认")
        self.click(BatchImportLocator.submit_btn)
        logger.info("已点击提交")

    @allure.step("断言样品存在于样品列表中")
    def assert_sample_exists(self, sample_name):
        """等待后调接口查询，断言样品存在，返回样品ID"""
        from object.sample_center.sample_list import SampleListPage
        self.page.wait_for_timeout(5000)
        sample_page = SampleListPage(self.page)
        sample_id = sample_page.get_sample_id(sample_name)
        assert sample_id is not None, f"未在样品列表中找到样品: {sample_name}"
        logger.info(f"样品已登记成功，样品ID: {sample_id}")
        return sample_id

    @staticmethod
    @allure.step("生成样品登记 Excel")
    def generate_sample_excel(project_name, sample_name, sample_form, batch_current_amount, sample_tube_type):
        """
        用 InTable 模板生成带数据的样品登记 Excel
        :param project_name: 所属项目
        :param sample_name: 样品名称
        :param sample_form: 样品形态
        :param batch_current_amount: 当前的量
        :param sample_tube_type: 样品管/孔板类型
        :return: 文件路径
        """
        wb = openpyxl.Workbook()
        ws = wb.active
        headers = [
            "样品ID", "样品名称", "关联的物质批号", "所属项目", "样品形态",
            "(溶液)摩尔浓度", "(溶液)溶剂", "当前的量", "分子量",
            "最末级库位ID", "盒内位置坐标", "样品管/孔板类型", "容器ID", "孔位",
            "生产日期", "失效日期", "条形码", "备注",
        ]
        for col_idx, h in enumerate(headers, 1):
            ws.cell(row=1, column=col_idx, value=h)

        data = {
            "样品ID": sample_name,
            "样品名称": sample_name,
            "所属项目": project_name,
            "样品形态": sample_form,
            "当前的量": batch_current_amount,
            "样品管/孔板类型": sample_tube_type,
            "容器ID": sample_name,
        }
        for col_idx, h in enumerate(headers, 1):
            val = data.get(h)
            if val is not None:
                ws.cell(row=2, column=col_idx, value=val)

        save_path = os.path.join(tempfile.gettempdir(), f"sample_register_{sample_name}.xlsx")
        wb.save(save_path)
        logger.info(f"样品登记 Excel 已生成: {save_path}")
        return save_path
